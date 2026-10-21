"""فواتير نقاط البيع بالساعة — عدد الفواتير وقيمتها في فترة تاريخ وساعات محددة.

قراءة فقط (SELECT) من رأس فاتورة نقاط البيع IAS_POS_BILL_MST:
- BILL_DATE تاريخ بلا وقت (كل الفواتير 00:00)، والوقت في BILL_TIME نصاً بصيغة HH24:MI:SS.
- القيمة = BILL_AMT + VAT_AMT (شاملة الضريبة) بنفس شروط تحليل المبيعات: غير معلّقة ومبلغ منطقي،
  فيطابق عدد الفواتير ما يظهر في تحليل المبيعات للفترة نفسها (قبل المرتجع).
- نطاق ساعات يعبر منتصف الليل (مثل 22:00 → 02:00) يُحسب ضمن تاريخ الفاتورة نفسه.
"""

from __future__ import annotations

import logging
import re
from datetime import date
from typing import Any

from django.core.cache import cache

from .oracle_stock import (
    OracleStockError,
    _as_date,
    _bill_amt_ok,
    _branch_names,
    _date_params,
    _fetch_all,
    _hung_ok,
    _norm_brn_code,
    _pos_owner,
    oracle_enabled,
)

logger = logging.getLogger(__name__)

_CACHE_TTL = 120
_CACHE_VER = "v4"
MAX_DAYS = 92
_TIME_RE = re.compile(r"^([01]\d|2[0-3]):([0-5]\d)$")


def parse_hhmm(value: Any, default: str) -> str:
    """«8:5» أو «08:05» أو «08:05:00» → «08:05»؛ قيمة غير صالحة ترفع ValueError."""
    text = str(value or "").strip()
    if not text:
        return default
    parts = text.split(":")
    if len(parts) not in (2, 3) or not all(p.isdigit() for p in parts):
        raise ValueError("وقت غير صالح — استخدم الصيغة ساعة:دقيقة مثل 08:30.")
    hhmm = f"{int(parts[0]):02d}:{int(parts[1]):02d}"
    if not _TIME_RE.match(hhmm):
        raise ValueError("وقت غير صالح — الساعة من 00 إلى 23 والدقائق من 00 إلى 59.")
    return hhmm


def normalize_window(time_from: str, time_to: str) -> tuple[str, str]:
    """نفس وقت البداية والنهاية (مثل 12:00 ص → 12:00 ص) يعني اليوم كاملاً لا دقيقة واحدة."""
    if time_from == time_to:
        return "00:00", "23:59"
    return time_from, time_to


def window_hours(time_from: str, time_to: str) -> list[int]:
    """الساعات التي يمسّها النطاق بالترتيب (يلتفّ بعد منتصف الليل)."""
    start, end = int(time_from[:2]), int(time_to[:2])
    if time_from <= time_to:
        return list(range(start, end + 1))
    # يعبر منتصف الليل؛ إن بدأ وانتهى في الساعة نفسها (22:30 → 22:10) لا تتكرر تلك الساعة
    return list(range(start, 24)) + [h for h in range(0, end + 1) if h < start]


def _time_clause(time_from: str, time_to: str) -> str:
    """شرط الوقت على BILL_TIME (نص HH24:MI:SS يُقارَن أبجدياً بشكل صحيح)؛ نهاية الدقيقة مشمولة."""
    if time_from == "00:00" and time_to == "23:59":
        return ""
    if time_from <= time_to:
        return "AND m.BILL_TIME >= :t_from AND m.BILL_TIME <= :t_to"
    return "AND (m.BILL_TIME >= :t_from OR m.BILL_TIME <= :t_to)"


def _f(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _money(value: float) -> str:
    return f"{value:,.2f}"


def _count(value: int) -> str:
    return f"{value:,}"


def _pct(part: float, whole: float) -> float:
    return round(part * 100.0 / whole, 1) if whole else 0.0


def _row(label: str, invoices: int, value: float, total_invoices: int, top_invoices: int) -> dict[str, Any]:
    share = _pct(invoices, total_invoices)
    return {
        "label": label,
        "invoices": invoices,
        "invoices_display": _count(invoices),
        "value": round(value, 2),
        "value_display": _money(value),
        "avg_display": _money(value / invoices) if invoices else "—",
        "share": share,
        "share_display": f"{share:.1f}%",
        # عرض شريط النسبة نسبةً لأعلى صف حتى يظهر الفرق بين الساعات بوضوح
        "bar_width": f"{(invoices * 100.0 / top_invoices) if top_invoices else 0:.1f}",
    }


def _per_day(invoices: int, days: int) -> str:
    """متوسط فواتير اليوم بمنزلة عشرية واحدة (12.5) ودون كسور حين يكون عدداً صحيحاً."""
    avg = invoices / days if days else 0.0
    return f"{avg:,.0f}" if float(avg).is_integer() else f"{avg:,.1f}"


def hour_label(hour: int, time_from: str = "00:00", time_to: str = "23:59") -> str:
    """عنوان صف الساعة بمداه الفعلي: الساعة الأولى والأخيرة قد تكونان جزئيتين.

    مثال: من 08:30 إلى 06:00 ← «08:30 – 08:59» … «06:00 – 06:00» بدل «06:00 – 06:59».
    """
    start_m, end_m = "00", "59"
    first, last = int(time_from[:2]), int(time_to[:2])
    # ساعة يبدأ وينتهي فيها نطاق يعبر منتصف الليل تحوي جزأين — تُعرض كاملة
    split = time_from > time_to and first == last
    if hour == first and not split:
        start_m = time_from[3:5]
    if hour == last and not split:
        end_m = time_to[3:5]
    return f"{hour:02d}:{start_m} – {hour:02d}:{end_m}"


def build_pos_hours_report(
    date_from,
    date_to,
    *,
    time_from: str = "00:00",
    time_to: str = "23:59",
    branch_code: str = "",
) -> dict[str, Any]:
    """عدد فواتير نقاط البيع وقيمتها في الفترة والساعات المحددة: بالساعة وبالفرع وباليوم."""
    if not oracle_enabled():
        raise OracleStockError("أوراكل غير مفعّل.")
    d_from, d_to = _as_date(date_from), _as_date(date_to)
    if d_from > d_to:
        raise OracleStockError("تاريخ البداية بعد تاريخ النهاية.")
    if (d_to - d_from).days + 1 > MAX_DAYS:
        raise OracleStockError(f"الفترة أطول من {MAX_DAYS} يوماً — اختر فترة أقصر.")
    time_from, time_to = normalize_window(parse_hhmm(time_from, "00:00"), parse_hhmm(time_to, "23:59"))
    brn = _norm_brn_code(branch_code)

    key = f"pos_hours:{_CACHE_VER}:{d_from}:{d_to}:{time_from}:{time_to}:{brn}"
    try:
        cached = cache.get(key)
    except Exception:  # noqa: BLE001
        cached = None
    if cached is not None:
        return cached

    pos = _pos_owner()
    params: dict[str, Any] = _date_params(d_from, d_to)
    time_sql = _time_clause(time_from, time_to)
    if time_sql:
        params["t_from"] = f"{time_from}:00"
        params["t_to"] = f"{time_to}:59"
    brn_sql = ""
    if brn:
        params["brn"] = int(brn) if brn.isdigit() else brn
        brn_sql = "AND m.BRN_NO = :brn"

    rows = _fetch_all(
        f"""
        SELECT
            TRUNC(m.BILL_DATE) AS D,
            m.BRN_NO AS BRN,
            SUBSTR(m.BILL_TIME, 1, 2) AS HH,
            COUNT(DISTINCT m.BILL_NO) AS INVOICES,
            ROUND(SUM(NVL(m.BILL_AMT, 0) + NVL(m.VAT_AMT, 0)), 2) AS VALUE_TOTAL
        FROM {pos}.IAS_POS_BILL_MST m
        WHERE m.BILL_DATE >= :d_from AND m.BILL_DATE < :d_to_excl
          AND {_hung_ok("m")}
          AND {_bill_amt_ok("m", "BILL_AMT")}
          {time_sql}
          {brn_sql}
        GROUP BY TRUNC(m.BILL_DATE), m.BRN_NO, SUBSTR(m.BILL_TIME, 1, 2)
        """,
        params,
    )

    names = _branch_names()
    by_hour: dict[int, list] = {}
    by_branch: dict[str, list] = {}
    by_day: dict[date, list] = {}
    hour_branch: dict[tuple[int, str], int] = {}
    day_branch: dict[tuple[date, str], int] = {}
    total_invoices = 0
    total_value = 0.0
    for r in rows:
        invoices = int(r.get("INVOICES") or 0)
        value = _f(r.get("VALUE_TOTAL"))
        try:
            hour = int(str(r.get("HH") or "").strip())
        except ValueError:
            continue  # BILL_TIME فارغ أو تالف — لا يقع في أي ساعة
        code = _norm_brn_code(r.get("BRN"))
        day = _as_date(r.get("D"))
        for bucket, k in ((by_hour, hour), (by_branch, code), (by_day, day)):
            acc = bucket.setdefault(k, [0, 0.0])
            acc[0] += invoices
            acc[1] += value
        hour_branch[(hour, code)] = hour_branch.get((hour, code), 0) + invoices
        day_branch[(day, code)] = day_branch.get((day, code), 0) + invoices
        total_invoices += invoices
        total_value += value

    # كل ساعات النطاق تظهر حتى بلا فواتير (صفر) حتى تُرى الساعات الهادئة
    hours = window_hours(time_from, time_to)
    top_hour = max((by_hour.get(h, [0])[0] for h in hours), default=0)
    hour_rows = [
        dict(_row(hour_label(h, time_from, time_to), *by_hour.get(h, [0, 0.0]), total_invoices, top_hour), hour=h)
        for h in hours
    ]
    busiest = max(hour_rows, key=lambda x: x["invoices"], default=None)

    # الأيام المقسوم عليها: أيام الفترة التي مضت فعلاً (فترة تنتهي في المستقبل لا تُقسَم على أيام لم تأتِ)
    days = (d_to - d_from).days + 1
    days_elapsed = max(1, (min(d_to, date.today()) - d_from).days + 1)

    top_branch = max((v[0] for v in by_branch.values()), default=0)
    branch_rows = sorted(
        (
            dict(_row(names.get(code, code or "—"), inv, val, total_invoices, top_branch), code=code)
            for code, (inv, val) in by_branch.items()
        ),
        key=lambda x: (-x["invoices"], x["code"]),
    )

    # عدد فواتير كل فرع في كل ساعة (أعمدة بأسماء الفروع، بترتيب الأكثر فواتير)
    branch_cols = [{"code": b["code"], "name": b["label"]} for b in branch_rows]
    for h in hour_rows:
        h["by_branch"] = [
            _count(hour_branch[(h["hour"], col["code"])]) if hour_branch.get((h["hour"], col["code"])) else "—"
            for col in branch_cols
        ]
    branch_totals = [b["invoices_display"] for b in branch_rows]
    # متوسط اليوم لكل فرع = إجمالي الفترة ÷ عدد الأيام
    for b in branch_rows:
        b["per_day_invoices_display"] = _per_day(b["invoices"], days_elapsed)
        b["per_day_value_display"] = _money(b["value"] / days_elapsed)
    branch_per_day = [b["per_day_invoices_display"] for b in branch_rows]

    day_rows = []
    if d_from != d_to:
        top_day = max((v[0] for v in by_day.values()), default=0)
        day_rows = [
            dict(
                _row(d.isoformat(), *by_day[d], total_invoices, top_day),
                date=d.isoformat(),
                by_branch=[
                    _count(day_branch[(d, col["code"])]) if day_branch.get((d, col["code"])) else "—"
                    for col in branch_cols
                ],
            )
            for d in sorted(by_day)
        ]
    report = {
        "from": d_from.isoformat(),
        "to": d_to.isoformat(),
        "days": days,
        "time_from": time_from,
        "time_to": time_to,
        "crosses_midnight": time_from > time_to,
        "branch": brn,
        "branch_name": names.get(brn, brn) if brn else "",
        "invoices": total_invoices,
        "invoices_display": _count(total_invoices),
        "value": round(total_value, 2),
        "value_display": _money(total_value),
        "avg_display": _money(total_value / total_invoices) if total_invoices else "—",
        "per_day_display": _count(round(total_invoices / days_elapsed)) if days > 1 else "",
        "days_elapsed": days_elapsed,
        "per_day_invoices_display": _per_day(total_invoices, days_elapsed),
        "per_day_value_display": _money(total_value / days_elapsed),
        "busiest_label": busiest["label"] if busiest and busiest["invoices"] else "—",
        "busiest_invoices_display": busiest["invoices_display"] if busiest and busiest["invoices"] else "",
        "hours": hour_rows,
        "branches": branch_rows,
        "branch_cols": branch_cols,
        "branch_totals": branch_totals,
        "branch_per_day": branch_per_day,
        "days_rows": day_rows,
    }
    try:
        cache.set(key, report, _CACHE_TTL)
    except Exception:  # noqa: BLE001
        pass
    return report


# ===== تصدير Excel: كل جدول في ورقة مستقلة بنفس تصميم الصفحة =====
# ألوان الأعمدة كما في الصفحة: الاسم · الأعداد · المبالغ · المتوسط · النسبة
_XL_COLORS = {
    "name": "DCE8F4",
    "count": "E6F0FA",
    "avg1": "E6F0FA",
    "money": "E8F4E3",
    "avg": "EEEAF8",
    "share": "FCEBD3",
    "head": "F6F6F9",
    "total": "EEF3F9",
    "perday": "E3ECF6",
}
_XL_FORMATS = {"count": "#,##0", "avg1": "#,##0.0", "money": "#,##0.00", "avg": "#,##0.00", "share": "0.0%"}


def _num(text: Any) -> float | int:
    """«1,234.5» أو «—» → رقم (الخلية الفارغة صفر)."""
    if isinstance(text, (int, float)):
        return text
    t = str(text or "").replace(",", "").strip()
    if not t or t == "—":
        return 0
    try:
        v = float(t)
    except ValueError:
        return 0
    return int(v) if v.is_integer() else v


def _cell(text: Any):
    """خانة فرع في جدول: «—» (لا فواتير) تبقى فارغة كما في الصفحة، وغيرها رقم."""
    return None if str(text or "").strip() in ("", "—") else _num(text)


def report_period_line(report: dict[str, Any]) -> str:
    """نفس سطر الفترة تحت عنوان كل جدول في الصفحة."""
    line = f"الفترة: من {report['from']}"
    if report["to"] != report["from"]:
        line += f" إلى {report['to']}"
    line += f" · الساعة {report['time_from']} – {report['time_to']}"
    if report.get("branch_name"):
        line += f" · {report['branch_name']}"
    return line


def _xl_sheet(wb, title: str, period: str, columns: list[tuple[str, str, int]], rows: list[list], footer: list[tuple[list, str]]):
    """ورقة: عنوان الجدول وسطر الفترة ثم الترويسة والصفوف وصفوف الإجمالي.

    columns: (العنوان، النوع، العرض) — النوع يحدد لون العمود وتنسيق الرقم.
    footer: [(القيم، لون الخلفية)] صفوف أسفل الجدول (الإجمالي / متوسط اليوم).
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title=title[:31])
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    n = len(columns)
    hair = Side(style="thin", color="D9DCE3")
    border = Border(left=hair, right=hair, top=hair, bottom=hair)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    font = "Arial"

    for row, text, size, color in ((1, title, 14, "1F2A37"), (2, period, 10, "5A6675")):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
        c = ws.cell(row=row, column=1, value=text)
        c.font = Font(name=font, size=size, bold=True, color=color)
        c.alignment = center
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18

    head_row = 4
    for i, (label, _kind, width) in enumerate(columns, start=1):
        c = ws.cell(row=head_row, column=i, value=label)
        c.font = Font(name=font, size=10, bold=True, color="5A6675")
        c.fill = PatternFill("solid", fgColor=_XL_COLORS["head"])
        c.alignment = center
        c.border = Border(left=hair, right=hair, top=hair, bottom=Side(style="medium", color="A8B0BD"))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[head_row].height = 22

    def write(r: int, values: list, fill: str | None = None, bold: bool = False):
        for i, ((_label, kind, _w), v) in enumerate(zip(columns, values), start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=font, size=10, bold=bold or kind == "name", color="1F2A37")
            c.alignment = center
            c.border = border
            color = fill or _XL_COLORS.get(kind)
            if color:
                c.fill = PatternFill("solid", fgColor=color)
            if kind in _XL_FORMATS and isinstance(v, (int, float)):
                # متوسط اليوم في أعمدة الأعداد قد يكون كسراً (466.5) — نُظهر منزلة عشرية بدل التقريب
                fractional = kind == "count" and isinstance(v, float) and not v.is_integer()
                c.number_format = "#,##0.0" if fractional else _XL_FORMATS[kind]

    r = head_row + 1
    for values in rows:
        write(r, values)
        r += 1
    for k, (values, fill) in enumerate(footer):
        write(r, values, fill=fill, bold=True)
        if k == 0:
            for i in range(1, n + 1):
                ws.cell(row=r, column=i).border = Border(
                    left=hair, right=hair, bottom=hair, top=Side(style="medium", color="4A7FB5"),
                )
        r += 1

    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.print_title_rows = f"{head_row}:{head_row}"
    # طباعة رسمية: رأسي بعرض الصفحة (مثل PDF)، بلا ترويسة ولا تذييل
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.page_margins.header = ws.page_margins.footer = 0
    return ws


def _avg(value: float, invoices: int) -> float:
    return round(value / invoices, 2) if invoices else 0


def build_pos_hours_excel(report: dict[str, Any]):
    """ملف xlsx: الفواتير بالساعة / بالفرع / باليوم — كل جدول في ورقة بنفس أعمدة الصفحة."""
    from io import BytesIO

    from django.http import HttpResponse
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    period = report_period_line(report)
    cols = report.get("branch_cols") or []
    multi = len(cols) > 1
    total_inv = report["invoices"] or 0
    share = (lambda inv: inv / total_inv if total_inv else 0)
    branch_heads = [(c["name"], "count", 12) for c in cols] if multi else []
    tail_heads = [("عدد الفواتير", "count", 13), ("القيمة", "money", 16), ("متوسط الفاتورة", "avg", 14), ("النسبة", "share", 10)]
    grand = [report["invoices"], report["value"], _avg(report["value"], report["invoices"]), 1 if total_inv else 0]
    branch_totals = [_num(x) for x in report.get("branch_totals", [])] if multi else []

    # 1) الفواتير بالساعة
    rows = [
        [h["label"]] + ([_cell(x) for x in h.get("by_branch", [])] if multi else [])
        + [h["invoices"], h["value"], _avg(h["value"], h["invoices"]), share(h["invoices"])]
        for h in report["hours"]
    ]
    _xl_sheet(
        wb, "الفواتير بالساعة", period, [("الساعة", "name", 16)] + branch_heads + tail_heads, rows,
        [(["الإجمالي"] + branch_totals + grand, _XL_COLORS["total"])],
    )

    # 2) الفواتير بالفرع
    daily = report.get("days", 1) > 1
    columns = [("#", "count", 6), ("الفرع", "name", 26), ("عدد الفواتير", "count", 13), ("القيمة", "money", 16)]
    if daily:
        columns += [("فواتير اليوم", "avg1", 12), ("قيمة اليوم", "money", 15)]
    columns += [("متوسط الفاتورة", "avg", 14), ("النسبة", "share", 10)]
    rows = []
    for i, b in enumerate(report["branches"], start=1):
        row = [i, b["label"], b["invoices"], b["value"]]
        if daily:
            row += [_num(b.get("per_day_invoices_display")), _num(b.get("per_day_value_display"))]
        rows.append(row + [_avg(b["value"], b["invoices"]), share(b["invoices"])])
    total = ["", "الإجمالي", report["invoices"], report["value"]]
    if daily:
        total += [_num(report.get("per_day_invoices_display")), _num(report.get("per_day_value_display"))]
    total += [_avg(report["value"], report["invoices"]), 1 if total_inv else 0]
    _xl_sheet(wb, "الفواتير بالفرع", period, columns, rows, [(total, _XL_COLORS["total"])] if rows else [])

    # 3) الفواتير باليوم (لفترة أكثر من يوم)
    if report.get("days_rows"):
        rows = [
            [d["label"]] + ([_cell(x) for x in d.get("by_branch", [])] if multi else [])
            + [d["invoices"], d["value"], _avg(d["value"], d["invoices"]), share(d["invoices"])]
            for d in report["days_rows"]
        ]
        per_day = (
            [f"متوسط اليوم (÷ {report.get('days_elapsed', 1)} يوم)"]
            + ([_num(x) for x in report.get("branch_per_day", [])] if multi else [])
            + [_num(report.get("per_day_invoices_display")), _num(report.get("per_day_value_display")), None, None]
        )
        _xl_sheet(
            wb, "الفواتير باليوم", period, [("التاريخ", "name", 14)] + branch_heads + tail_heads, rows,
            [(["الإجمالي"] + branch_totals + grand, _XL_COLORS["total"]), (per_day, _XL_COLORS["perday"])],
        )

    buf = BytesIO()
    wb.save(buf)
    stamp = f"{report['from']}_{report['to']}".replace("-", "")
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="pos_invoices_by_hour_{stamp}.xlsx"'
    return response
