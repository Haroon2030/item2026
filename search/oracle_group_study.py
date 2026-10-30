"""دراسة مجموعة أصناف — لوحة تحليل مالية وتشغيلية ومبيعاتية وتجارية وتوريدية لفترة وفرع.

قراءة فقط (SELECT). المصادر:
- المبيعات: نقاط البيع (IAS_POS_BILL_*) + نظام المبيعات النقدي المباشر 1/5 (بلا خلاصات ترحيل نقاط البيع)
  + الآجل 4/8 (IAS_BILL_*)، ومرتجعاتها (IAS_POS_RT_BILL_* / IAS_RT_BILL_* بلا خلاصات مرتجعات نقاط البيع).
  المبالغ بدون ضريبة كباقي شاشات المبيعات.
- الوارد: سندات التوريد (IAS_GR_INCOMING_*) + فواتير الشراء (IAS_PI_BILL_*) ومرتجعاتها (IAS_PR_BILL_*).
  التحويلات بين المخازن ليست وارداً.
- أوامر الصرف (IAS_OUTGOING_*): التوالف (41101006 توالف، 41101008 الأصناف المفقودة والتالفة)،
  والمرتجع للمورد (21101001 موردين) منفصلاً، والباقي «صرف آخر».
- الرصيد الحالي (IAS_ITM_WCODE) لمخازن الفرع المختار.
الكميات بالوحدة الرئيسية للصنف (MAIN_UNIT في أونكس: كيلو للموزون، علبة/برطمان/حبة/تنك…)؛ تكلفة المبيعات والرصيد = الكمية × متوسط التكلفة الحالي للصنف (كما في شاشة تحليل المبيعات)،
والصرف (توالف/مورد) بتكلفة أمر الصرف نفسه (كما في تقرير أوامر الصرف)، والبيع بسعر البيع.
"""

from __future__ import annotations

import logging
from datetime import date
from typing import Any

from django.core.cache import cache

from .oracle_item_branch_report import _bill_net_sql
from .oracle_stock import (
    OracleStockError,
    _as_date,
    _bind_gcode,
    _branch_names,
    _date_params,
    _fetch_all,
    _norm_brn_code,
    _pos_dtl_net_sql,
    _pos_owner,
    _run_parallel,
    _schema,
    fetch_item_max_pack_map,
    fetch_warehouse_options,
    oracle_enabled,
    oracle_session,
)

logger = logging.getLogger(__name__)

_CACHE_TTL = 600
_CACHE_VER = "v17"
MAX_DAYS = 400

WASTE_ACCOUNTS = ("41101006", "41101008")  # توالف · الأصناف المفقودة والتالفة
SUPPLIER_RETURN_ACCOUNTS = ("21101001",)  # موردين — صرف بضاعة للمورد
POS_SALES_SUMMARY = "خلاصة مبيعات رقم نقطة البيع%"
POS_RETURNS_SUMMARY = "رقم نقطة البيع =%"
STOCK_QTY = "NVL(d.P_QTY, NVL(d.I_QTY, 0) * NVL(NULLIF(d.P_SIZE, 0), 1))"
GROUP_SQL = "i.G_CODE = :g"  # شرط الأصناف؛ «كمية الطلب» تمرّر بدله «كل المجموعات عدا…»

_AR_MONTHS = ("يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
              "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر")


def _f(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _money(value: float) -> str:
    return f"{value:,.2f}"


def _qty(value: float) -> str:
    v = round(float(value or 0), 2)
    return f"{v:,.0f}" if v.is_integer() else f"{v:,.2f}"


def _pct(part: float, whole: float) -> float | None:
    return round(part * 100.0 / whole, 1) if whole else None


def _pct_display(value: float | None) -> str:
    return "—" if value is None else f"{value:.1f}%"


def _code(value: Any) -> str:
    text = str(value or "").strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


# ===== الاستعلامات =====

def _brn_filter(alias: str, brn: str, params: dict) -> str:
    if not brn:
        return ""
    params["brn"] = int(brn) if brn.isdigit() else brn
    return f"AND {alias}.BRN_NO = :brn"


def _pos_sales(params: dict, brn_sql: str, grp: str = GROUP_SQL) -> list[dict]:
    pos, sch = _pos_owner(), _schema()
    return _fetch_all(
        f"""
        SELECT TO_CHAR(d.I_CODE) AS I, TRUNC(m.BILL_DATE, 'MM') AS M, TO_CHAR(m.BRN_NO) AS B,
               SUM({STOCK_QTY}) AS Q, ROUND(SUM({_pos_dtl_net_sql("d")}), 2) AS V,
               COUNT(DISTINCT m.BRN_NO || '-' || m.BILL_NO) AS N
        FROM {pos}.IAS_POS_BILL_DTL d
        JOIN {pos}.IAS_POS_BILL_MST m
          ON m.BILL_NO = d.BILL_NO AND m.BRN_NO = d.BRN_NO AND NVL(m.BILL_SRL, 0) = NVL(d.BILL_SRL, 0)
        JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
        WHERE m.BILL_DATE >= :d_from AND m.BILL_DATE < :d_to_excl
          AND (m.HUNG IS NULL OR m.HUNG = 0) AND {grp} {brn_sql.replace('{a}', 'm')}
        GROUP BY TO_CHAR(d.I_CODE), TRUNC(m.BILL_DATE, 'MM'), TO_CHAR(m.BRN_NO)
        """,
        params,
    )


def _bill_sales(params: dict, brn_sql: str, doc_types: tuple, *, exclude_summary: bool, grp: str = GROUP_SQL) -> list[dict]:
    sch = _schema()
    types = ", ".join(str(int(t)) for t in doc_types)
    summ = f"AND NVL(b.C_NAME, ' ') NOT LIKE '{POS_SALES_SUMMARY}'" if exclude_summary else ""
    return _fetch_all(
        f"""
        SELECT TO_CHAR(d.I_CODE) AS I, TRUNC(b.BILL_DATE, 'MM') AS M, TO_CHAR(b.BRN_NO) AS B,
               SUM({STOCK_QTY}) AS Q, ROUND(SUM({_bill_net_sql("d")}), 2) AS V,
               COUNT(DISTINCT b.BILL_SER) AS N, MAX(NVL(d.P_SIZE, 1)) AS PS
        FROM {sch}.IAS_BILL_DTL d
        JOIN {sch}.IAS_BILL_MST b
          ON b.BILL_NO = d.BILL_NO AND b.BILL_SER = d.BILL_SER AND b.BILL_DOC_TYPE = d.BILL_DOC_TYPE
        JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
        WHERE b.BILL_DATE >= :d_from AND b.BILL_DATE < :d_to_excl
          AND (b.CNCL_FLG IS NULL OR b.CNCL_FLG = 0) AND ABS(NVL(b.BILL_AMT, 0)) <= 5000000
          AND b.BILL_DOC_TYPE IN ({types}) {summ}
          AND {grp} {brn_sql.replace('{a}', 'b')}
        GROUP BY TO_CHAR(d.I_CODE), TRUNC(b.BILL_DATE, 'MM'), TO_CHAR(b.BRN_NO)
        """,
        params,
    )


def _pos_returns(params: dict, brn_sql: str, grp: str = GROUP_SQL) -> list[dict]:
    pos, sch = _pos_owner(), _schema()
    return _fetch_all(
        f"""
        SELECT TO_CHAR(d.I_CODE) AS I, TO_CHAR(m.BRN_NO) AS B,
               SUM({STOCK_QTY}) AS Q, ROUND(SUM({_pos_dtl_net_sql("d")}), 2) AS V
        FROM {pos}.IAS_POS_RT_BILL_DTL d
        JOIN {pos}.IAS_POS_RT_BILL_MST m ON m.RT_BILL_NO = d.RT_BILL_NO AND m.BRN_NO = d.BRN_NO
        JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
        WHERE m.RT_BILL_DATE >= :d_from AND m.RT_BILL_DATE < :d_to_excl
          AND (m.HUNG IS NULL OR m.HUNG = 0) AND {grp} {brn_sql.replace('{a}', 'm')}
        GROUP BY TO_CHAR(d.I_CODE), TO_CHAR(m.BRN_NO)
        """,
        params,
    )


def _bill_returns(params: dict, brn_sql: str, doc_types: tuple, *, exclude_summary: bool, grp: str = GROUP_SQL) -> list[dict]:
    sch = _schema()
    types = ", ".join(str(int(t)) for t in doc_types)
    summ = f"AND NVL(r.C_NAME, ' ') NOT LIKE '{POS_RETURNS_SUMMARY}'" if exclude_summary else ""
    return _fetch_all(
        f"""
        SELECT TO_CHAR(d.I_CODE) AS I, TO_CHAR(r.BRN_NO) AS B,
               SUM({STOCK_QTY}) AS Q, ROUND(SUM({_bill_net_sql("d")}), 2) AS V
        FROM {sch}.IAS_RT_BILL_DTL d
        JOIN {sch}.IAS_RT_BILL_MST r ON r.RT_BILL_SER = d.RT_BILL_SER AND r.BRN_NO = d.BRN_NO
        JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
        WHERE r.RT_BILL_DATE >= :d_from AND r.RT_BILL_DATE < :d_to_excl
          AND (r.CNCL_FLG IS NULL OR r.CNCL_FLG = 0) AND ABS(NVL(r.BILL_AMT, 0)) <= 5000000
          AND r.RT_BILL_DOC_TYPE IN ({types}) {summ}
          AND {grp} {brn_sql.replace('{a}', 'r')}
        GROUP BY TO_CHAR(d.I_CODE), TO_CHAR(r.BRN_NO)
        """,
        params,
    )


def _supply_vouchers(params: dict, brn_sql: str) -> list[dict]:
    sch = _schema()
    return _fetch_all(
        f"""
        SELECT TO_CHAR(d.I_CODE) AS I, TO_CHAR(m.BRN_NO) AS B,
               SUM({STOCK_QTY}) AS Q, COUNT(DISTINCT m.INC_SER) AS N, MAX(m.INC_DATE) AS LAST_D
        FROM {sch}.IAS_GR_INCOMING_DTL d
        JOIN {sch}.IAS_GR_INCOMING_MST m ON m.INC_SER = d.INC_SER AND m.INCOM_TYPE = d.INCOM_TYPE
        JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
        WHERE m.INC_DATE >= :d_from AND m.INC_DATE < :d_to_excl
          AND i.G_CODE = :g {brn_sql.replace('{a}', 'm')}
        GROUP BY TO_CHAR(d.I_CODE), TO_CHAR(m.BRN_NO)
        """,
        params,
    )


def _purchases(params: dict, brn_sql: str) -> list[dict]:
    sch = _schema()
    return _fetch_all(
        f"""
        SELECT TO_CHAR(d.I_CODE) AS I, TO_CHAR(m.BRN_NO) AS B, SUM({STOCK_QTY}) AS Q,
               ROUND(SUM(NVL(d.I_QTY, 0) * NVL(d.I_PRICE, 0) - NVL(d.DIS_AMT, 0)), 2) AS V,
               COUNT(DISTINCT m.BILL_SER) AS N, COUNT(DISTINCT m.V_CODE) AS VENDORS, MAX(m.BILL_DATE) AS LAST_D
        FROM {sch}.IAS_PI_BILL_DTL d
        JOIN {sch}.IAS_PI_BILL_MST m
          ON m.BILL_NO = d.BILL_NO AND m.BILL_SER = d.BILL_SER AND m.BILL_DOC_TYPE = d.BILL_DOC_TYPE
        JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
        WHERE m.BILL_DATE >= :d_from AND m.BILL_DATE < :d_to_excl
          AND NVL(m.HUNG, 0) = 0 AND i.G_CODE = :g {brn_sql.replace('{a}', 'm')}
        GROUP BY TO_CHAR(d.I_CODE), TO_CHAR(m.BRN_NO)
        """,
        params,
    )


def _purchase_returns(params: dict, brn_sql: str) -> list[dict]:
    sch = _schema()
    return _fetch_all(
        f"""
        SELECT TO_CHAR(d.I_CODE) AS I, TO_CHAR(r.BRN_NO) AS B, SUM({STOCK_QTY}) AS Q,
               ROUND(SUM(NVL(d.I_QTY, 0) * NVL(d.I_PRICE, 0) - NVL(d.DIS_AMT, 0)), 2) AS V
        FROM {sch}.IAS_PR_BILL_DTL d
        JOIN {sch}.IAS_PR_BILL_MST r
          ON r.RT_BILL_NO = d.RT_BILL_NO AND r.RT_BILL_SER = d.RT_BILL_SER AND r.RT_BILL_DOC_TYPE = d.RT_BILL_DOC_TYPE
        JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
        WHERE r.RT_BILL_DATE >= :d_from AND r.RT_BILL_DATE < :d_to_excl
          AND NVL(r.HUNG, 0) = 0 AND i.G_CODE = :g {brn_sql.replace('{a}', 'r')}
        GROUP BY TO_CHAR(d.I_CODE), TO_CHAR(r.BRN_NO)
        """,
        params,
    )


def _issues(params: dict, brn_sql: str) -> list[dict]:
    sch = _schema()
    return _fetch_all(
        f"""
        SELECT TO_CHAR(d.I_CODE) AS I, TO_CHAR(m.A_CODE) AS A, TO_CHAR(m.BRN_NO) AS B,
               SUM({STOCK_QTY}) AS Q, ROUND(SUM(NVL(d.I_QTY, 0) * NVL(d.STK_COST, 0)), 2) AS C,
               COUNT(DISTINCT m.OUT_SER) AS N
        FROM {sch}.IAS_OUTGOING_DTL d
        JOIN {sch}.IAS_OUTGOING_MST m ON m.OUT_SER = d.OUT_SER AND m.OUT_TYPE = d.OUT_TYPE
        JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
        WHERE m.OUT_TYPE = 1 AND m.OUT_DATE >= :d_from AND m.OUT_DATE < :d_to_excl
          AND (m.HUNG IS NULL OR m.HUNG = 0) AND i.G_CODE = :g {brn_sql.replace('{a}', 'm')}
        GROUP BY TO_CHAR(d.I_CODE), TO_CHAR(m.A_CODE), TO_CHAR(m.BRN_NO)
        """,
        params,
    )


def _distinct_counts(params: dict, brn_sql: str) -> dict[str, int]:
    """عدد المستندات المميّزة للمجموعة: فاتورة فيها 3 أصناف من المجموعة تُعدّ مرة واحدة."""
    pos, sch = _pos_owner(), _schema()
    waste = ", ".join(f"'{a}'" for a in WASTE_ACCOUNTS)
    vret = ", ".join(f"'{a}'" for a in SUPPLIER_RETURN_ACCOUNTS)
    row = _fetch_all(
        f"""
        SELECT
          (SELECT COUNT(DISTINCT m.BRN_NO || '-' || m.BILL_NO) FROM {pos}.IAS_POS_BILL_DTL d
             JOIN {pos}.IAS_POS_BILL_MST m ON m.BILL_NO = d.BILL_NO AND m.BRN_NO = d.BRN_NO AND NVL(m.BILL_SRL, 0) = NVL(d.BILL_SRL, 0)
             JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
            WHERE m.BILL_DATE >= :d_from AND m.BILL_DATE < :d_to_excl AND (m.HUNG IS NULL OR m.HUNG = 0)
              AND i.G_CODE = :g {brn_sql.replace('{a}', 'm')}) AS POS_N,
          (SELECT COUNT(DISTINCT b.BILL_SER || '-' || b.BILL_DOC_TYPE) FROM {sch}.IAS_BILL_DTL d
             JOIN {sch}.IAS_BILL_MST b ON b.BILL_NO = d.BILL_NO AND b.BILL_SER = d.BILL_SER AND b.BILL_DOC_TYPE = d.BILL_DOC_TYPE
             JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
            WHERE b.BILL_DATE >= :d_from AND b.BILL_DATE < :d_to_excl AND (b.CNCL_FLG IS NULL OR b.CNCL_FLG = 0)
              AND ABS(NVL(b.BILL_AMT, 0)) <= 5000000 AND b.BILL_DOC_TYPE IN (1, 4, 5, 8)
              AND NVL(b.C_NAME, ' ') NOT LIKE '{POS_SALES_SUMMARY}'
              AND i.G_CODE = :g {brn_sql.replace('{a}', 'b')}) AS BILL_N,
          (SELECT COUNT(DISTINCT m.INC_SER || '-' || m.INCOM_TYPE) FROM {sch}.IAS_GR_INCOMING_DTL d
             JOIN {sch}.IAS_GR_INCOMING_MST m ON m.INC_SER = d.INC_SER AND m.INCOM_TYPE = d.INCOM_TYPE
             JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
            WHERE m.INC_DATE >= :d_from AND m.INC_DATE < :d_to_excl AND i.G_CODE = :g {brn_sql.replace('{a}', 'm')}) AS SUP_N,
          (SELECT COUNT(DISTINCT m.BILL_SER || '-' || m.BILL_DOC_TYPE) FROM {sch}.IAS_PI_BILL_DTL d
             JOIN {sch}.IAS_PI_BILL_MST m ON m.BILL_NO = d.BILL_NO AND m.BILL_SER = d.BILL_SER AND m.BILL_DOC_TYPE = d.BILL_DOC_TYPE
             JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
            WHERE m.BILL_DATE >= :d_from AND m.BILL_DATE < :d_to_excl AND NVL(m.HUNG, 0) = 0
              AND i.G_CODE = :g {brn_sql.replace('{a}', 'm')}) AS PI_N,
          (SELECT COUNT(DISTINCT m.OUT_SER) FROM {sch}.IAS_OUTGOING_DTL d
             JOIN {sch}.IAS_OUTGOING_MST m ON m.OUT_SER = d.OUT_SER AND m.OUT_TYPE = d.OUT_TYPE
             JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
            WHERE m.OUT_TYPE = 1 AND m.OUT_DATE >= :d_from AND m.OUT_DATE < :d_to_excl AND (m.HUNG IS NULL OR m.HUNG = 0)
              AND TO_CHAR(m.A_CODE) IN ({waste}) AND i.G_CODE = :g {brn_sql.replace('{a}', 'm')}) AS WASTE_N,
          (SELECT COUNT(DISTINCT m.OUT_SER) FROM {sch}.IAS_OUTGOING_DTL d
             JOIN {sch}.IAS_OUTGOING_MST m ON m.OUT_SER = d.OUT_SER AND m.OUT_TYPE = d.OUT_TYPE
             JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
            WHERE m.OUT_TYPE = 1 AND m.OUT_DATE >= :d_from AND m.OUT_DATE < :d_to_excl AND (m.HUNG IS NULL OR m.HUNG = 0)
              AND TO_CHAR(m.A_CODE) IN ({vret}) AND i.G_CODE = :g {brn_sql.replace('{a}', 'm')}) AS VRET_N
        FROM DUAL
        """,
        params,
    )[0]
    return {k.lower(): int(_f(v)) for k, v in row.items()}


def _branch_invoices(params: dict, brn_sql: str) -> dict[str, int]:
    """عدد فواتير البيع المميّزة لكل فرع (نقاط البيع + النقدي المباشر والآجل بلا خلاصات الترحيل)."""
    pos, sch = _pos_owner(), _schema()
    rows = _fetch_all(
        f"""
        SELECT B, SUM(N) AS N FROM (
          SELECT TO_CHAR(m.BRN_NO) AS B, COUNT(DISTINCT m.BILL_NO) AS N
          FROM {pos}.IAS_POS_BILL_DTL d
          JOIN {pos}.IAS_POS_BILL_MST m ON m.BILL_NO = d.BILL_NO AND m.BRN_NO = d.BRN_NO AND NVL(m.BILL_SRL, 0) = NVL(d.BILL_SRL, 0)
          JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
          WHERE m.BILL_DATE >= :d_from AND m.BILL_DATE < :d_to_excl AND (m.HUNG IS NULL OR m.HUNG = 0)
            AND i.G_CODE = :g {brn_sql.replace('{a}', 'm')}
          GROUP BY TO_CHAR(m.BRN_NO)
          UNION ALL
          SELECT TO_CHAR(b.BRN_NO) AS B, COUNT(DISTINCT b.BILL_SER || '-' || b.BILL_DOC_TYPE) AS N
          FROM {sch}.IAS_BILL_DTL d
          JOIN {sch}.IAS_BILL_MST b ON b.BILL_NO = d.BILL_NO AND b.BILL_SER = d.BILL_SER AND b.BILL_DOC_TYPE = d.BILL_DOC_TYPE
          JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
          WHERE b.BILL_DATE >= :d_from AND b.BILL_DATE < :d_to_excl AND (b.CNCL_FLG IS NULL OR b.CNCL_FLG = 0)
            AND ABS(NVL(b.BILL_AMT, 0)) <= 5000000 AND b.BILL_DOC_TYPE IN (1, 4, 5, 8)
            AND NVL(b.C_NAME, ' ') NOT LIKE '{POS_SALES_SUMMARY}'
            AND i.G_CODE = :g {brn_sql.replace('{a}', 'b')}
          GROUP BY TO_CHAR(b.BRN_NO)
        ) GROUP BY B
        """,
        params,
    )
    return {_norm_brn_code(r["B"]): int(_f(r["N"])) for r in rows}


def _items_and_stock(group: Any, brn: str, sub: str = "") -> tuple[dict[str, dict], dict[str, tuple[float, float]]]:
    """أصناف المجموعة (الاسم) + رصيدها الحالي وقيمته بمتوسط التكلفة لمخازن الفرع (أو كل المخازن)."""
    sch = _schema()
    sub_sql = "AND TO_CHAR(i.MNG_CODE) = :mng" if sub else ""
    base: dict[str, Any] = {"g": group} | ({"mng": sub} if sub else {})
    items = {
        _code(r["I"]): {"name": str(r.get("N") or "").strip()}
        for r in _fetch_all(f"SELECT TO_CHAR(i.I_CODE) AS I, i.I_NAME AS N FROM {sch}.IAS_ITM_MST i WHERE i.G_CODE = :g {sub_sql}", base)
    }
    params: dict[str, Any] = dict(base)
    wh_sql = ""
    if brn:
        # مخازن الفرع من قائمة المخازن المعتمدة في التطبيق (جدول المخازن نفسه بلا عمود فرع)
        whs = [
            _code(w.get("code")) for w in fetch_warehouse_options(active_only=False) or []
            if _norm_brn_code(w.get("branch_code")) == brn and _code(w.get("code"))
        ]
        if not whs:
            return items, {}
        for n, wh in enumerate(whs):
            params[f"w{n}"] = int(wh) if wh.isdigit() else wh
        wh_sql = f"AND w.W_CODE IN ({', '.join(f':w{n}' for n in range(len(whs)))})"
    stock = {
        _code(r["I"]): (_f(r["Q"]), _f(r["UC"]))
        for r in _fetch_all(
            f"""
            SELECT TO_CHAR(w.I_CODE) AS I, SUM(NVL(w.AVL_QTY, 0)) AS Q,
                   AVG(CASE WHEN NVL(w.I_CWTAVG, w.PRIMARY_COST) > 0 THEN NVL(w.I_CWTAVG, w.PRIMARY_COST) END) AS UC
            FROM {sch}.IAS_ITM_WCODE w
            JOIN {sch}.IAS_ITM_MST i ON i.I_CODE = w.I_CODE
            WHERE i.G_CODE = :g {sub_sql} {wh_sql}
            GROUP BY TO_CHAR(w.I_CODE)
            """,
            params,
        )
    }
    return items, stock


def _group_name(group: Any) -> str:
    rows = _fetch_all(f"SELECT G_A_NAME AS N FROM {_schema()}.GROUP_DETAILS WHERE G_CODE = :g", {"g": group})
    return str((rows[0].get("N") if rows else "") or "").strip()


def _subgroup_name(group: Any, sub: str) -> str:
    rows = _fetch_all(
        f"SELECT MNG_A_NAME AS N FROM {_schema()}.IAS_MAINSUB_GRP_DTL WHERE G_CODE = :g AND TO_CHAR(MNG_CODE) = :mng",
        {"g": group, "mng": sub},
    )
    return str((rows[0].get("N") if rows else "") or "").strip()


def _base_units(group: Any, sub: str = "", grp: str = GROUP_SQL) -> dict[str, str]:
    """الوحدة الرئيسية لكل صنف (MAIN_UNIT = 1 في أونكس، حجمها 1): كيلو للموزون، علبة، حبة، تنك…
    وإن لم تُعلَّم وحدة رئيسية فأصغر وحدة."""
    rows = _fetch_all(
        f"""
        SELECT I, U FROM (
            SELECT TO_CHAR(d.I_CODE) AS I, TRIM(d.ITM_UNT) AS U,
                   ROW_NUMBER() OVER (
                       PARTITION BY d.I_CODE
                       ORDER BY CASE WHEN d.MAIN_UNIT = 1 THEN 0 ELSE 1 END, NVL(d.P_SIZE, 0), d.ITM_UNT
                   ) AS RN
            FROM {_schema()}.IAS_ITM_DTL d
            JOIN {_schema()}.IAS_ITM_MST i ON i.I_CODE = d.I_CODE
            WHERE {grp} AND NVL(d.P_SIZE, 0) > 0 {"AND TO_CHAR(i.MNG_CODE) = :mng" if sub else ""}
        ) WHERE RN = 1
        """,
        ({"g": group} if ":g" in grp else {}) | ({"mng": sub} if sub else {}),
    )
    return {_code(r["I"]): str(r.get("U") or "").replace("ـ", "").strip() for r in rows}


def _pack_label(pack: float, unit: str, base: str) -> str:
    """«كرتون (12 حبة)»: الوحدة الكبرى ثم ما تحويه من الوحدة الصغرى."""
    unit = unit.replace("ـ", "").strip().rstrip("0123456789").strip()  # «كرتون1» في الكرت = كرتون
    if pack <= 1:
        return unit or base or "—"
    inner = f"{_qty(pack)} {base}".strip()
    return f"{unit} ({inner})" if unit else inner


# ===== التجميع =====

def build_group_study(date_from, date_to, *, group_code: str, branch_code: str = "", subgroup_code: str = "") -> dict[str, Any]:
    """دراسة المجموعة: مؤشرات، جدول الأصناف، اتجاه شهري، وقراءات (مالية، مبيعاتية، تشغيلية، تجارية، توريدية)."""
    if not oracle_enabled():
        raise OracleStockError("أوراكل غير مفعّل.")
    gcode = str(group_code or "").strip()
    if not gcode:
        raise OracleStockError("اختر المجموعة.")
    d_from, d_to = _as_date(date_from), _as_date(date_to)
    if d_from > d_to:
        raise OracleStockError("تاريخ البداية بعد تاريخ النهاية.")
    days = (d_to - d_from).days + 1
    if days > MAX_DAYS:
        raise OracleStockError(f"الفترة أطول من {MAX_DAYS} يوماً — اختر فترة أقصر.")
    days_elapsed = max(1, (min(d_to, date.today()) - d_from).days + 1)
    brn = _norm_brn_code(branch_code)
    sub = str(subgroup_code or "").strip()

    key = f"group_study:{_CACHE_VER}:{d_from}:{d_to}:{gcode}:{sub}:{brn}"
    try:
        cached = cache.get(key)
    except Exception:  # noqa: BLE001
        cached = None
    if cached is not None:
        return cached

    group = _bind_gcode(gcode)
    params: dict[str, Any] = dict(_date_params(d_from, d_to), g=group)
    brn_sql = _brn_filter("{a}", brn, params)
    if sub:  # كل استعلامات الدراسة تربط الصنف بالاسم المستعار i
        params["mng"] = sub
        brn_sql += " AND TO_CHAR(i.MNG_CODE) = :mng"

    def job(fn, *args, **kwargs):
        def run():
            with oracle_session():
                return fn(*args, **kwargs)
        return run

    (pos_rows, cash_rows, credit_rows, pos_ret, cash_ret, credit_ret,
     supply, pi, pr, issues, counts, items_stock, brn_inv) = _run_parallel(
        [
            job(_pos_sales, params, brn_sql),
            job(_bill_sales, params, brn_sql, (1, 5), exclude_summary=True),
            job(_bill_sales, params, brn_sql, (4, 8), exclude_summary=False),
            job(_pos_returns, params, brn_sql),
            job(_bill_returns, params, brn_sql, (1, 5), exclude_summary=True),
            job(_bill_returns, params, brn_sql, (4, 8), exclude_summary=False),
            job(_supply_vouchers, params, brn_sql),
            job(_purchases, params, brn_sql),
            job(_purchase_returns, params, brn_sql),
            job(_issues, params, brn_sql),
            job(_distinct_counts, params, brn_sql),
            job(_items_and_stock, group, brn, sub),
            job(_branch_invoices, params, brn_sql),
        ],
        max_workers=6,
        timeout_sec=240,
    )
    items, stock = items_stock
    group_name = _group_name(group) or gcode
    names = _branch_names()

    def blank() -> dict[str, float]:
        return {k: 0.0 for k in (
            "pos_q", "pos_v", "cash_q", "cash_v", "credit_q", "credit_v", "inv",
            "ret_q", "ret_v", "sup_q", "sup_n", "pi_q", "pi_v", "pi_n", "pr_q", "pr_v",
            "waste_q", "waste_c", "vret_q", "vret_c", "other_q", "other_c",
        )}

    per: dict[str, dict[str, float]] = {}
    # (الفرع، الصنف) ← الكميات والقيم، لجدول الفروع: التكلفة = صافي كمية الصنف في الفرع × متوسط تكلفته
    per_brn: dict[tuple[str, str], dict[str, float]] = {}
    monthly: dict[date, dict[str, float]] = {}
    month_q: dict[tuple[date, str], float] = {}

    def acc(code: str) -> dict[str, float]:
        return per.setdefault(code, blank())

    def bacc(r: dict) -> dict[str, float]:
        return per_brn.setdefault((_norm_brn_code(r.get("B")) or brn, _code(r["I"])), blank())

    for rows, qk, vk in ((pos_rows, "pos_q", "pos_v"), (cash_rows, "cash_q", "cash_v"), (credit_rows, "credit_q", "credit_v")):
        for r in rows:
            for a in (acc(_code(r["I"])), bacc(r)):
                a[qk] += _f(r["Q"])
                a[vk] += _f(r["V"])
                a["inv"] += _f(r["N"])
            m = _as_date(r["M"])
            mb = monthly.setdefault(m, {"v": 0.0, "q": 0.0})
            mb["v"] += _f(r["V"])
            month_q[(m, _code(r["I"]))] = month_q.get((m, _code(r["I"])), 0.0) + _f(r["Q"])
    for rows in (pos_ret, cash_ret, credit_ret):
        for r in rows:
            for a in (acc(_code(r["I"])), bacc(r)):
                a["ret_q"] += _f(r["Q"])
                a["ret_v"] += _f(r["V"])
    for r in supply:
        for a in (acc(_code(r["I"])), bacc(r)):
            a["sup_q"] += _f(r["Q"])
            a["sup_n"] += _f(r["N"])
    for r in pi:
        for a in (acc(_code(r["I"])), bacc(r)):
            a["pi_q"] += _f(r["Q"])
            a["pi_v"] += _f(r["V"])
            a["pi_n"] += _f(r["N"])
    for r in pr:
        for a in (acc(_code(r["I"])), bacc(r)):
            a["pr_q"] += _f(r["Q"])
            a["pr_v"] += _f(r["V"])
    for r in issues:
        acct = str(r.get("A") or "").strip()
        kind = "waste" if acct in WASTE_ACCOUNTS else "vret" if acct in SUPPLIER_RETURN_ACCOUNTS else "other"
        for a in (acc(_code(r["I"])), bacc(r)):
            a[f"{kind}_q"] += _f(r["Q"])
            a[f"{kind}_c"] += _f(r["C"])
    for code in stock:
        acc(code)

    codes = sorted(per)
    packs = fetch_item_max_pack_map(codes) if codes else {}
    bases = _base_units(gcode, sub) if codes else {}

    # الكميات بالوحدة الرئيسية للصنف (وحدة المخزون، حجمها 1): كيلو للموزون، علبة، حبة…
    # فلا تحويل؛ العبوة الكبرى («كرتون (11 كيلو)») تُعرض تلميحاً فقط
    for (m, code), q in month_q.items():
        monthly[m]["q"] += q

    rows_out: list[dict[str, Any]] = []
    T = blank() | {"cost": 0.0, "stock_q": 0.0, "stock_v": 0.0}
    for code in codes:
        a = per[code]
        sold_q = a["pos_q"] + a["cash_q"] + a["credit_q"]
        sold_v = a["pos_v"] + a["cash_v"] + a["credit_v"]
        net_q = sold_q - a["ret_q"]
        net_v = sold_v - a["ret_v"]
        st_q, unit_cost = stock.get(code, (0.0, 0.0))
        cost = net_q * unit_cost  # متوسط التكلفة الحالي للصنف — كما في شاشة تحليل المبيعات
        # الصرف (توالف/مورد/أخرى) يبقى بتكلفة أمر الصرف نفسه (الكمية × STK_COST يوم الصرف)
        # كما في تقرير أوامر الصرف والقيد المحاسبي، لا بمتوسط التكلفة الحالي
        gross = net_v - cost
        incoming_q = a["sup_q"] + a["pi_q"] - a["pr_q"]
        daily = net_q / days_elapsed if days_elapsed else 0.0
        cover = (st_q / daily) if daily > 0 else None
        pk = packs.get(code) or {}
        row = {
            "code": code,
            "name": (items.get(code) or {}).get("name") or code,
            "pack": _qty(_f(pk.get("pack"))) if pk.get("pack") else "—",
            "unit": str(pk.get("unit") or "").strip(),
            "base_unit": bases.get(code, ""),
            # الوحدة الرئيسية هي ما يُعرض (الكميات بها)؛ العبوة الكبرى «كرتون (11 كيلو)» تلميح فقط
            "pack_label": bases.get(code) or "—",
            "pack_hint": _pack_label(_f(pk.get("pack")), str(pk.get("unit") or ""), bases.get(code, ""))
            if _f(pk.get("pack")) > 1 else "",
            "invoices": int(a["inv"]),
            # قبل المرتجع = رقم تقرير أونكس «المبيعات حسب الأصناف»؛ الصافي بعد خصم المرتجع
            "sold_q": sold_q, "sold_v": sold_v, "ret_q": a["ret_q"],
            "net_q": net_q, "net_v": net_v, "cost": cost, "gross": gross,
            "price": (net_v / net_q) if net_q > 0 else None,  # متوسط سعر بيع الوحدة الرئيسية بدون ضريبة
            "margin": _pct(gross, net_v),
            "ret_v": a["ret_v"],
            "incoming_q": incoming_q, "sup_q": a["sup_q"], "pi_q": a["pi_q"], "pi_v": a["pi_v"],
            "waste_q": a["waste_q"], "waste_c": a["waste_c"], "vret_q": a["vret_q"],
            "stock_q": st_q, "stock_v": st_q * unit_cost, "unit_cost": unit_cost,
            "cover_days": cover,
            "sell_through": _pct(net_q, net_q + st_q) if (net_q + st_q) > 0 else None,
        }
        rows_out.append(row)
        for k in a:
            T[k] += a[k]
        T["cost"] += cost
        T["stock_q"] += st_q
        T["stock_v"] += st_q * unit_cost

    for r in rows_out:
        r.update({
            "net_q_display": _qty(r["net_q"]), "net_v_display": _money(r["net_v"]),
            "sold_q_display": _qty(r["sold_q"]), "sold_v_display": _money(r["sold_v"]),
            "ret_q_display": _qty(r["ret_q"]), "ret_v_display": _money(r["ret_v"]),
            "cost_display": _money(r["cost"]), "gross_display": _money(r["gross"]),
            "margin_display": _pct_display(r["margin"]), "invoices_display": f"{r['invoices']:,}",
            "incoming_q_display": _qty(r["incoming_q"]), "waste_q_display": _qty(r["waste_q"]),
            "waste_c_display": _money(r["waste_c"]), "vret_q_display": _qty(r["vret_q"]),
            "stock_q_display": _qty(r["stock_q"]), "stock_v_display": _money(r["stock_v"]),
            "unit_cost_display": _money(r["unit_cost"]),
            "price_display": "—" if r["price"] is None else _money(r["price"]),
            "cover_display": "—" if r["cover_days"] is None else f"{r['cover_days']:,.0f} يوم",
            "sell_through_display": _pct_display(r["sell_through"]),
            "share": _pct(r["net_v"], T["pos_v"] + T["cash_v"] + T["credit_v"] - T["ret_v"]),
        })
        r["share_display"] = _pct_display(r["share"])
    rows_out.sort(key=lambda r: (-r["net_v"], r["code"]))

    sold_v = T["pos_v"] + T["cash_v"] + T["credit_v"]
    net_v = sold_v - T["ret_v"]
    net_q = T["pos_q"] + T["cash_q"] + T["credit_q"] - T["ret_q"]
    gross = net_v - T["cost"]
    margin = _pct(gross, net_v)
    incoming_q = T["sup_q"] + T["pi_q"] - T["pr_q"]
    daily_q = net_q / days_elapsed if days_elapsed else 0.0
    cover = (T["stock_q"] / daily_q) if daily_q > 0 else None
    active = [r for r in rows_out if r["net_q"] > 0]

    kpis = {
        "net_v": net_v, "net_v_display": _money(net_v),
        "sold_v_display": _money(sold_v), "ret_v_display": _money(T["ret_v"]),
        "sold_q_display": _qty(T["pos_q"] + T["cash_q"] + T["credit_q"]), "ret_q_display": _qty(T["ret_q"]),
        "ret_pct_display": _pct_display(_pct(T["ret_v"], sold_v)),
        "net_q_display": _qty(net_q), "invoices_display": f"{counts['pos_n'] + counts['bill_n']:,}",
        "per_day_display": _money(net_v / days_elapsed),
        "cost_display": _money(T["cost"]), "gross": gross, "gross_display": _money(gross),
        "margin_display": _pct_display(margin),
        "pos_v_display": _money(T["pos_v"]), "cash_v_display": _money(T["cash_v"]), "credit_v_display": _money(T["credit_v"]),
        "pos_share_display": _pct_display(_pct(T["pos_v"], sold_v)),
        "incoming_q_display": _qty(incoming_q), "sup_q_display": _qty(T["sup_q"]), "sup_n": counts["sup_n"],
        "pi_q_display": _qty(T["pi_q"]), "pi_v_display": _money(T["pi_v"]), "pi_n": counts["pi_n"],
        "pr_v_display": _money(T["pr_v"]),
        "waste_q_display": _qty(T["waste_q"]), "waste_c_display": _money(T["waste_c"]),
        "waste_pct_display": _pct_display(_pct(T["waste_c"], T["cost"])),
        "waste_orders": counts["waste_n"],
        "vret_q_display": _qty(T["vret_q"]), "vret_c_display": _money(T["vret_c"]), "vret_orders": counts["vret_n"],
        "stock_q_display": _qty(T["stock_q"]), "stock_v_display": _money(T["stock_v"]),
        "cover_display": "—" if cover is None else f"{cover:,.0f} يوم",
        "items_total": len(items), "items_active": len(active),
    }

    # الاتجاه الشهري
    peak = max((m["v"] for m in monthly.values()), default=0.0)
    trend = [
        {
            "label": f"{_AR_MONTHS[m.month - 1]} {m.year}", "v_display": _money(v["v"]), "q_display": _qty(v["q"]),
            "bar": f"{(v['v'] * 100.0 / peak) if peak else 0:.1f}",
        }
        for m, v in sorted(monthly.items())
    ]

    branches = _branch_rows(per_brn, stock, brn_inv, names, net_v)
    readings = _readings(rows_out, kpis, T, net_v, sold_v, net_q, gross, margin, cover, incoming_q, days_elapsed)
    assessment = _assessment(rows_out, kpis, T, sold_v, net_q, margin, cover, incoming_q)

    report = {
        "from": d_from.isoformat(), "to": d_to.isoformat(), "days": days, "days_elapsed": days_elapsed,
        "group_code": gcode, "group_name": group_name,
        "subgroup_code": sub, "subgroup_name": (_subgroup_name(group, sub) or sub) if sub else "",
        "branch": brn, "branch_name": names.get(brn, brn) if brn else "",
        "kpis": kpis, "rows": rows_out, "branches": branches, "trend": trend, "readings": readings,
        "assessment": assessment,
        "has_purchase_invoices": counts["pi_n"] > 0,
    }
    try:
        cache.set(key, report, _CACHE_TTL)
    except Exception:  # noqa: BLE001
        pass
    return report


def _branch_rows(per_brn, stock, brn_inv, names, total_net_v) -> dict[str, Any]:
    """تفصيل كل فرع: قنوات البيع والمرتجع والصافي والتكلفة والربح والوارد والصرف، مع سطر الإجمالي.
    بلا رصيد للفروع: تكلفة الصنف واحدة في كل الفروع، ولا يُعرض رصيد بمتوسط مخازن الفرع."""
    agg: dict[str, dict[str, float]] = {}
    keys = ("pos_v", "cash_v", "credit_v", "ret_v", "net_q", "net_v", "cost", "incoming_q",
            "waste_q", "waste_c", "vret_q", "vret_c")

    def slot(b: str) -> dict[str, float]:
        return agg.setdefault(b, {k: 0.0 for k in keys})

    for (b, code), a in per_brn.items():
        unit_cost = stock.get(code, (0.0, 0.0))[1]
        net_q = a["pos_q"] + a["cash_q"] + a["credit_q"] - a["ret_q"]
        s = slot(b)
        s["pos_v"] += a["pos_v"]
        s["cash_v"] += a["cash_v"]
        s["credit_v"] += a["credit_v"]
        s["ret_v"] += a["ret_v"]
        s["net_q"] += net_q
        s["net_v"] += a["pos_v"] + a["cash_v"] + a["credit_v"] - a["ret_v"]
        s["cost"] += net_q * unit_cost
        s["incoming_q"] += a["sup_q"] + a["pi_q"] - a["pr_q"]
        for k in ("waste_q", "waste_c", "vret_q", "vret_c"):
            s[k] += a[k]

    def finish(label: str, s: dict[str, float], invoices: int) -> dict[str, Any]:
        gross = s["net_v"] - s["cost"]
        margin = _pct(gross, s["net_v"])
        share = _pct(s["net_v"], total_net_v)
        return s | {
            "name": label, "invoices": invoices, "gross": gross, "margin": margin, "share": share,
            "invoices_display": f"{invoices:,}", "gross_display": _money(gross),
            "margin_display": _pct_display(margin), "share_display": _pct_display(share),
            "net_q_display": _qty(s["net_q"]), "incoming_q_display": _qty(s["incoming_q"]),
            "waste_q_display": _qty(s["waste_q"]), "vret_q_display": _qty(s["vret_q"]),
        } | {f"{k}_display": _money(s[k]) for k in ("pos_v", "cash_v", "credit_v", "ret_v", "net_v", "cost", "waste_c", "vret_c")}

    rows = [
        finish(names.get(b, b), s, brn_inv.get(b, 0)) | {"code": b}
        for b, s in agg.items()
        if any(abs(v) > 0.005 for v in s.values())
    ]
    rows.sort(key=lambda r: (-r["net_v"], r["code"]))
    total = {k: sum(r[k] for r in rows) for k in keys}
    return {"rows": rows, "total": finish("الإجمالي", total, sum(r["invoices"] for r in rows))}


def _assessment(rows, kpis, T, sold_v, net_q, margin, cover, incoming_q) -> dict[str, list[tuple[str, str, str]]]:
    """ملخص التحليل: نقاط القوة ونقاط تحتاج متابعة — كل نقطة (البند، القيمة، المعيار) بمعايير ثابتة معلنة."""
    good: list[tuple[str, str, str]] = []
    watch: list[tuple[str, str, str]] = []
    top = [r for r in rows if r["net_v"] > 0]

    if margin is not None:
        if margin >= 15:
            good.append(("نسبة مجمل الربح", kpis["margin_display"], "15% فأكثر"))
        elif margin < 10:
            watch.append(("نسبة مجمل الربح", kpis["margin_display"], "أقل من 10%"))
    neg = [r for r in top if r["gross"] < 0]
    low = [r for r in top if r["margin"] is not None and 0 <= r["margin"] < 5]
    if top and not neg:
        good.append(("أصناف بمجمل ربح سالب", "لا يوجد", "صفر"))
    if neg:
        watch.append(("أصناف بمجمل ربح سالب", f"{len(neg)} صنف، أكبرها «{neg[0]['name']}»", "يجب أن تكون صفراً"))
    if low:
        watch.append(("أصناف نسبة مجمل ربحها أقل من 5%", f"{len(low)} صنف", "أقل من 5% للصنف"))

    ret_pct = _pct(T["ret_v"], sold_v)
    if ret_pct is not None:
        if ret_pct < 2:
            good.append(("نسبة المرتجعات إلى إجمالي المبيعات", kpis["ret_pct_display"], "أقل من 2%"))
        elif ret_pct >= 3:
            watch.append(("نسبة المرتجعات إلى إجمالي المبيعات", kpis["ret_pct_display"], "3% فأكثر"))

    waste_pct = _pct(T["waste_c"], T["cost"])
    if waste_pct is not None:
        if waste_pct < 1:
            good.append(("نسبة التوالف إلى تكلفة المبيعات", kpis["waste_pct_display"], "أقل من 1%"))
        else:
            watch.append(("نسبة التوالف إلى تكلفة المبيعات", f"{kpis['waste_pct_display']} ({kpis['waste_c_display']})", "1% فأكثر"))

    if cover is not None:
        if 7 <= cover <= 60:
            good.append(("مدة تغطية المخزون", kpis["cover_display"], "بين 7 و60 يوماً"))
        elif cover > 90:
            watch.append(("مدة تغطية المخزون", kpis["cover_display"], "أكثر من 90 يوماً"))
        elif cover < 7:
            watch.append(("مدة تغطية المخزون", kpis["cover_display"], "أقل من 7 أيام"))
    dead = [r for r in rows if r["stock_q"] > 0 and r["net_q"] <= 0]
    if dead:
        watch.append(("أصناف لها رصيد دون حركة بيع", f"{len(dead)} صنف بتكلفة {_money(sum(r['stock_v'] for r in dead))}", "صفر"))
    short = [r for r in rows if r["cover_days"] is not None and r["cover_days"] < 7 and r["net_q"] > 0]
    if short:
        watch.append(("أصناف رصيدها يغطي أقل من 7 أيام", f"{len(short)} صنف", "7 أيام فأكثر"))

    total, active = kpis["items_total"], kpis["items_active"]
    if total:
        if active >= total * 0.7:
            good.append(("الأصناف المباعة من أصناف المجموعة", f"{active} من {total}", "70% فأكثر"))
        elif total - active > total * 0.3:
            watch.append(("أصناف دون مبيعات خلال الفترة", f"{total - active} من {total}", "أكثر من 30% من الأصناف"))

    if net_q > 0:
        cov = incoming_q * 100.0 / net_q
        if 90 <= cov <= 120:
            good.append(("الوارد مقابل الكمية المباعة", f"{cov:.0f}%", "بين 90% و120%"))
        elif cov < 90:
            watch.append(("الوارد مقابل الكمية المباعة", f"{cov:.0f}%", "أقل من 90% (سحب من المخزون)"))
        else:
            watch.append(("الوارد مقابل الكمية المباعة", f"{cov:.0f}%", "أكثر من 120% (تراكم مخزون)"))
    return {"good": good, "watch": watch}


def _readings(rows, kpis, T, net_v, sold_v, net_q, gross, margin, cover, incoming_q, days_elapsed) -> list[dict]:
    """القراءات بصياغة محاسبية: بند وقيمته من أرقام الدراسة نفسها — كل قراءة: عنوان، بنود، ومستوى (ok / watch / risk)."""
    out: list[dict] = []
    top = [r for r in rows if r["net_v"] > 0]

    # المؤشرات المالية
    fin = [
        f"صافي المبيعات: {kpis['net_v_display']}",
        f"تكلفة المبيعات (بمتوسط التكلفة الحالي): {kpis['cost_display']}",
        f"مجمل الربح: {kpis['gross_display']}",
        f"نسبة مجمل الربح: {kpis['margin_display']}",
    ]
    neg = [r for r in top if r["gross"] < 0]
    low = [r for r in top if r["margin"] is not None and 0 <= r["margin"] < 5]
    if neg:
        fin.append(f"أصناف بمجمل ربح سالب: {len(neg)} صنف، أكبرها «{neg[0]['name']}» بخسارة {_money(-neg[0]['gross'])}")
    if low:
        fin.append(f"أصناف نسبة مجمل ربحها أقل من 5%: {len(low)} صنف")
    by_profit = sorted(top, key=lambda r: -r["gross"])[:5]
    if gross > 0 and by_profit:
        share = sum(r["gross"] for r in by_profit) * 100.0 / gross
        fin.append(f"مساهمة أعلى 5 أصناف في مجمل الربح: {share:.0f}%")
    out.append({"title": "المؤشرات المالية", "level": "risk" if (margin is not None and margin < 5) or neg else "ok", "points": fin})

    # مؤشرات المبيعات
    sales = [
        f"عدد الفواتير: {kpis['invoices_display']}",
        f"الكمية المباعة الصافية: {kpis['net_q_display']} وحدة",
        f"متوسط المبيعات اليومية: {kpis['per_day_display']} (عن {days_elapsed} يوماً)",
    ]
    if net_v > 0 and top:
        run, n80 = 0.0, 0
        for r in top:
            run += r["net_v"]
            n80 += 1
            if run >= 0.8 * net_v:
                break
        sales.append(f"أصناف تحقق 80% من صافي المبيعات: {n80} من {len(top)} صنف مباع")
        sales.append(f"نسبة مبيعات نقاط البيع: {kpis['pos_share_display']}")
        sales.append(f"نسبة المرتجعات إلى إجمالي المبيعات: {kpis['ret_pct_display']}")
    out.append({"title": "مؤشرات المبيعات", "level": "ok", "points": sales})

    # المخزون والصرف
    ops = [
        f"رصيد المخزون الحالي: {kpis['stock_q_display']} وحدة بتكلفة {kpis['stock_v_display']}",
        f"مدة تغطية المخزون: {kpis['cover_display']} (بمتوسط البيع اليومي)",
    ]
    if T["waste_q"]:
        ops.append(f"التوالف: {kpis['waste_q_display']} وحدة بتكلفة {kpis['waste_c_display']} ({kpis['waste_pct_display']} من تكلفة المبيعات) في {kpis['waste_orders']} أمر صرف")
    if T["vret_q"]:
        ops.append(f"المصروف للمورد: {kpis['vret_q_display']} وحدة بتكلفة {kpis['vret_c_display']} في {kpis['vret_orders']} أمر صرف")
    dead = [r for r in rows if r["stock_q"] > 0 and r["net_q"] <= 0]
    if dead:
        dead_v = sum(r["stock_v"] for r in dead)
        ops.append(f"أصناف لها رصيد دون حركة بيع خلال الفترة: {len(dead)} صنف بتكلفة {_money(dead_v)}")
    short = [r for r in rows if r["cover_days"] is not None and r["cover_days"] < 7 and r["net_q"] > 0]
    if short:
        ops.append(f"أصناف رصيدها يغطي أقل من 7 أيام: {len(short)} صنف، منها «{short[0]['name']}»")
    level = "risk" if (cover is not None and cover > 90) or dead else "watch" if short else "ok"
    out.append({"title": "المخزون والصرف", "level": level, "points": ops})

    # تشكيلة الأصناف
    com = [f"الأصناف المباعة خلال الفترة: {kpis['items_active']} من {kpis['items_total']} صنف في المجموعة"]
    if top:
        best = top[0]
        com.append(f"الصنف الأعلى مبيعاً: «{best['name']}»")
        com.append(f"صافي مبيعاته: {best['net_v_display']} ({best['share_display']} من المجموعة) بنسبة مجمل ربح {best['margin_display']}")
    idle = kpis["items_total"] - kpis["items_active"]
    if idle > 0:
        com.append(f"أصناف دون مبيعات خلال الفترة: {idle} صنف")
    out.append({"title": "تشكيلة الأصناف", "level": "watch" if idle > kpis["items_total"] * 0.3 else "ok", "points": com})

    # المشتريات والتوريد
    sup = []
    if kpis["pi_n"]:
        sup.append(f"فواتير الشراء: {kpis['pi_n']:,} فاتورة بكمية {kpis['pi_q_display']} وقيمة {kpis['pi_v_display']}")
        sup.append(f"مرتجعات الشراء: {kpis['pr_v_display']}")
    else:
        sup.append("فواتير الشراء: لا توجد فواتير شراء لأصناف المجموعة خلال الفترة")
    if T["sup_q"]:
        sup.append(f"سندات التوريد: {kpis['sup_n']} سند بكمية {kpis['sup_q_display']} وحدة")
    if net_q > 0:
        cov = incoming_q * 100.0 / net_q
        gap = "الفرق مسحوب من المخزون" if cov < 90 else "متوازن مع المبيعات" if cov <= 120 else "الزيادة مضافة إلى المخزون"
        sup.append(f"الوارد مقابل المباع: {kpis['incoming_q_display']} مقابل {kpis['net_q_display']} وحدة ({cov:.0f}%)، {gap}")
    out.append({"title": "المشتريات والتوريد", "level": "watch" if not kpis["pi_n"] else "ok", "points": sup})
    return out


# ===== تصدير Excel: ملخص وقراءات وأصناف وشهري وقنوات — كل جدول في ورقة بتنسيق الشاشة =====

def _num_text(text: Any) -> float:
    return _f(str(text or "").replace(",", ""))


def build_group_study_excel(report: dict[str, Any]):
    """ملف xlsx منظم للدراسة: الملخص، القراءات، تحليل الأصناف، المبيعات الشهرية، قنوات البيع."""
    from io import BytesIO

    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from .oracle_pos_hours import _XL_COLORS, _xl_sheet

    k = report["kpis"]
    sub = f" · فرعية {report['subgroup_code']} - {report['subgroup_name']}" if report.get("subgroup_code") else ""
    scope = f"{report['group_code']} - {report['group_name']}{sub} · {report.get('branch_name') or 'كل الفروع'}"
    period = f"{scope} · من {report['from']} إلى {report['to']} ({report['days_elapsed']} يوم) · بدون ضريبة"
    wb = Workbook()
    wb.remove(wb.active)
    hair = Side(style="thin", color="D9DCE3")
    border = Border(left=hair, right=hair, top=hair, bottom=hair)
    right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # 1) الملخص: المؤشر والقيمة والتفصيل
    summary = [
        ["صافي المبيعات", k["net_v_display"], f"قبل المرتجع {k['sold_v_display']} · المرتجع {k['ret_v_display']} ({k['ret_pct_display']})"],
        ["الربح الإجمالي", k["gross_display"], f"الهامش {k['margin_display']} · التكلفة {k['cost_display']}"],
        ["الفواتير", k["invoices_display"], f"{k['net_q_display']} وحدة · نقاط البيع {k['pos_share_display']}"],
        ["قنوات البيع", "", f"نقاط البيع {k['pos_v_display']} · نقدي مباشر {k['cash_v_display']} · آجل {k['credit_v_display']}"],
        ["الوارد (وحدة)", k["incoming_q_display"], f"سندات توريد {k['sup_n']} ({k['sup_q_display']}) · فواتير شراء {k['pi_n']} ({k['pi_v_display']})"],
        ["التوالف", k["waste_c_display"], f"{k['waste_q_display']} وحدة · {k['waste_orders']} أمر صرف · {k['waste_pct_display']} من التكلفة"],
        ["المصروف للمورد", k["vret_c_display"], f"{k['vret_q_display']} وحدة · {k['vret_orders']} أمر صرف"],
        ["الرصيد الحالي", k["stock_v_display"], f"{k['stock_q_display']} وحدة · يغطي {k['cover_display']}"],
        ["الأصناف", f"{k['items_active']} / {k['items_total']}", "صنف نشط من أصناف المجموعة"],
    ]
    ws = _xl_sheet(wb, "الملخص", period, [("المؤشر", "name", 22), ("القيمة", "text", 22), ("التفصيل", "text", 80)], summary, [])
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        row[2].alignment = right

    # 2) القراءات: عنوان كل قراءة ومستواها ثم نقاطها
    levels = {"risk": ("تحتاج إجراء", "FBE7E7"), "watch": ("للمتابعة", "FCEBD3"), "ok": ("جيد", "E3F4EC")}
    ws = wb.create_sheet("القراءات")
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 110
    for cell_ref, text, size, color in (("A1", "قراءات الدراسة", 14, "1F2A37"), ("A2", period, 10, "5A6675")):
        ws.merge_cells(f"{cell_ref}:C{cell_ref[1]}")
        ws[cell_ref] = text
        ws[cell_ref].font = Font(name="Arial", size=size, bold=True, color=color)
        ws[cell_ref].alignment = Alignment(horizontal="center")
    r = 4
    for g in report["readings"]:
        label, color = levels.get(g["level"], ("", "EEF3F9"))
        for i, point in enumerate(g["points"]):
            ws.cell(row=r, column=1, value=g["title"] if i == 0 else None)
            ws.cell(row=r, column=2, value=label if i == 0 else None)
            ws.cell(row=r, column=3, value=point)
            for c in range(1, 4):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                cell.alignment = right
                cell.font = Font(name="Arial", size=10, bold=c < 3, color="1F2A37")
                if c < 3:
                    cell.fill = PatternFill("solid", fgColor=color)
            r += 1
        r += 1
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # 3) تحليل الأصناف: نفس أعمدة الشاشة بأرقام قابلة للحساب
    columns = [
        ("الكود", "name", 18), ("الصنف", "name", 36), ("الوحدة الرئيسية", "text", 14), ("الفواتير", "count", 11),
        ("الكمية قبل المرتجع", "count", 13), ("المبيعات قبل المرتجع", "money", 15), ("كمية المرتجع", "count", 11),
        ("قيمة المرتجع", "money", 12), ("الكمية المباعة", "count", 13), ("سعر البيع", "avg", 12), ("صافي المبيعات", "money", 15), ("الحصة", "share", 9),
        ("تكلفة الوحدة", "avg", 12), ("التكلفة", "money", 15), ("الوارد", "count", 11), ("الرصيد", "count", 11),
    ]
    items = report["rows"]
    rows = []
    for x in items:
        rows.append([
            x["code"], x["name"], x["pack_label"], x["invoices"],
            round(x["sold_q"], 2), round(x["sold_v"], 2), round(x["ret_q"], 2), round(x["ret_v"], 2),
            round(x["net_q"], 2), None if x["price"] is None else round(x["price"], 2), round(x["net_v"], 2), (x["share"] or 0) / 100.0, round(x["unit_cost"], 2),
            round(x["cost"], 2), round(x["incoming_q"], 2), round(x["stock_q"], 2),
        ])
    net_v = sum(x["net_v"] for x in items)
    total = [
        "", "الإجمالي", "", None,
        round(sum(x["sold_q"] for x in items), 2), round(sum(x["sold_v"] for x in items), 2),
        round(sum(x["ret_q"] for x in items), 2), round(sum(x["ret_v"] for x in items), 2),
        round(sum(x["net_q"] for x in items), 2), None, round(net_v, 2), 1 if net_v else 0,
        None, round(sum(x["cost"] for x in items), 2),
        round(sum(x["incoming_q"] for x in items), 2), round(sum(x["stock_q"] for x in items), 2),
    ]
    ws = _xl_sheet(wb, "تحليل الأصناف", period, columns, rows, [(total, _XL_COLORS["total"])])
    ws.page_setup.orientation = "landscape"

    # 4) تفصيل الفروع
    br = report.get("branches") or {"rows": [], "total": None}
    if br["rows"]:
        def brow(x: dict, code: str) -> list:
            return [
                code, x["name"], x["invoices"], round(x["pos_v"], 2), round(x["cash_v"], 2), round(x["credit_v"], 2),
                round(x["ret_v"], 2), round(x["net_v"], 2), (x["share"] or 0) / 100.0, round(x["net_q"], 2),
                round(x["cost"], 2), round(x["incoming_q"], 2),
            ]

        ws = _xl_sheet(
            wb, "الفروع", period,
            [("الرقم", "name", 9), ("الفرع", "name", 26), ("الفواتير", "count", 11), ("نقاط البيع", "money", 15),
             ("نقدي مباشر", "money", 14), ("آجل", "money", 12), ("المرتجع", "money", 12), ("صافي المبيعات", "money", 15),
             ("الحصة", "share", 9), ("الكمية", "count", 12), ("التكلفة", "money", 15), ("الوارد", "count", 11)],
            [brow(x, x["code"]) for x in br["rows"]],
            [(brow(br["total"], ""), _XL_COLORS["total"])],
        )
        ws.page_setup.orientation = "landscape"

    # 5) المبيعات الشهرية
    _xl_sheet(
        wb, "المبيعات الشهرية", period,
        [("الشهر", "name", 18), ("الكمية", "count", 14), ("المبيعات", "money", 18)],
        [[m["label"], _num_text(m["q_display"]), _num_text(m["v_display"])] for m in report["trend"]],
        [],
    )

    # 6) قنوات البيع
    _xl_sheet(
        wb, "قنوات البيع", period,
        [("القناة", "name", 20), ("المبيعات", "money", 20)],
        [["نقاط البيع", _num_text(k["pos_v_display"])], ["نقدي مباشر", _num_text(k["cash_v_display"])],
         ["آجل", _num_text(k["credit_v_display"])], ["المرتجعات", -_num_text(k["ret_v_display"])]],
        [(["الصافي", _num_text(k["net_v_display"])], _XL_COLORS["total"])],
    )

    # 7) ملخص التحليل: نقاط القوة ونقاط تحتاج متابعة — البند والقيمة والمعيار
    a = report.get("assessment") or {"good": [], "watch": []}
    ws = wb.create_sheet("ملخص التحليل")
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    for col, width in (("A", 6), ("B", 40), ("C", 34), ("D", 30)):
        ws.column_dimensions[col].width = width
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row, text, size, color in ((1, "ملخص التحليل", 14, "1F2A37"), (2, period, 10, "5A6675")):
        ws.merge_cells(f"A{row}:D{row}")
        ws[f"A{row}"] = text
        ws[f"A{row}"].font = Font(name="Arial", size=size, bold=True, color=color)
        ws[f"A{row}"].alignment = center
    ws.merge_cells("A3:D3")
    ws["A3"] = f"نقاط القوة: {len(a['good'])} · نقاط تحتاج متابعة: {len(a['watch'])}"
    ws["A3"].font = Font(name="Arial", size=10, bold=True, color="1F2A37")
    ws["A3"].alignment = center
    r = 5
    for title, items, color, empty in (
        ("نقاط القوة", a["good"], "E3F4EC", "لا توجد بنود تحقق معايير القوة."),
        ("نقاط تحتاج متابعة", a["watch"], "FCEBD3", "لا توجد بنود تحتاج متابعة."),
    ):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = ws.cell(row=r, column=1, value=title)
        c.font = Font(name="Arial", size=12, bold=True, color="1F2A37")
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = right
        r += 1
        for i, label in enumerate(("م", "البند", "القيمة", "المعيار"), start=1):
            c = ws.cell(row=r, column=i, value=label)
            c.font = Font(name="Arial", size=10, bold=True, color="5A6675")
            c.fill = PatternFill("solid", fgColor=_XL_COLORS["head"])
            c.alignment = center
            c.border = border
        r += 1
        if not items:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            ws.cell(row=r, column=1, value=empty).alignment = right
            r += 1
        for n, (label, value, rule) in enumerate(items, start=1):
            for i, v in enumerate((n, label, value, rule), start=1):
                c = ws.cell(row=r, column=i, value=v)
                c.font = Font(name="Arial", size=10, bold=i == 2, color="1F2A37")
                c.alignment = center if i in (1, 3) else right
                c.border = border
            r += 1
        r += 1
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    stamp = f"{report['from']}_{report['to']}".replace("-", "")
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    code = f"{report['group_code']}_{report['subgroup_code']}" if report.get("subgroup_code") else report["group_code"]
    response["Content-Disposition"] = f'attachment; filename="group_study_{code}_{stamp}.xlsx"'
    return response
